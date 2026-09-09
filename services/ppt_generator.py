"""Builds a .pptx file from deck + slides data (as stored in MySQL)."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Simple palette — swap for brand colors if needed
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT = RGBColor(0x2F, 0x6F, 0xED)
DARK_TEXT = RGBColor(0x22, 0x22, 0x22)
LIGHT_GRAY = RGBColor(0xF4, 0xF6, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def _add_rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _add_text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return box


def _title_slide(prs, startup_name, tagline):
    slide = _blank_slide(prs)
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    _add_text(slide, Inches(1), Inches(2.9), Inches(11.3), Inches(1.2),
               startup_name, 44, WHITE, bold=True, align=PP_ALIGN.LEFT)
    _add_text(slide, Inches(1), Inches(4.0), Inches(11.3), Inches(0.8),
               tagline, 20, RGBColor(0xC9, 0xD6, 0xF0), align=PP_ALIGN.LEFT)


def _content_slide(prs, title, bullets, stats=None):
    slide = _blank_slide(prs)
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    _add_text(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(0.9),
               title, 30, NAVY, bold=True)

    body_top = Inches(1.6)
    body_left = Inches(0.7)
    body_width = Inches(11.9) if not stats else Inches(7.2)

    box = slide.shapes.add_textbox(body_left, body_top, body_width, Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {bullet}"
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(14)

    if stats:
        card = _add_rect(slide, Inches(8.3), body_top, Inches(4.3), Inches(4.6), LIGHT_GRAY)
        stat_box = slide.shapes.add_textbox(Inches(8.6), Inches(1.85), Inches(3.7), Inches(4.1))
        stf = stat_box.text_frame
        stf.word_wrap = True
        for i, stat in enumerate(stats):
            p = stf.paragraphs[0] if i == 0 else stf.add_paragraph()
            p.text = stat
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = ACCENT
            p.space_after = Pt(12)


def build_pptx(deck, output_path):
    """
    deck: dict from models.db.get_deck_with_slides(), i.e.
      {"startup_name": ..., "slides": [{"title", "slide_type", "content": {...}}, ...]}
    Writes the file to output_path and returns it.
    """
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    tagline = deck.get("idea_description", "")[:120]
    _title_slide(prs, deck["startup_name"], tagline)

    for slide_data in deck["slides"]:
        content = slide_data.get("content", {})
        bullets = content.get("bullets", [])
        stats = content.get("stats")
        _content_slide(prs, slide_data["title"], bullets, stats)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    return output_path
